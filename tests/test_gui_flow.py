"""Pruebas del flujo completo: validar credenciales, Siguiente y navegacion."""

from __future__ import annotations

import os
import shutil
import tempfile
import threading
import time
import unittest
from dataclasses import replace
from unittest.mock import patch

try:
    import tkinter as tk

    from sipclient import storage
    from sipclient.call import CallSession, CallState
    from sipclient.config import AccountConfig
    from sipclient.datastore import (
        DIRECTION_INCOMING,
        DIRECTION_OUTGOING,
        STATUS_COMPLETED,
        STATUS_NO_ANSWER,
        STATUS_REJECTED,
    )
    from sipclient.registrar import RegistrarState
    from sipclient.tonos import DummyTones
    from sipclient.ui import dialpad_page
    from sipclient.ui.app import MainWindow
    from sipclient.ui.common import SIDEBAR_ACTIVE
except Exception:  # pragma: no cover - sin tkinter disponible
    tk = None  # type: ignore[assignment]

from .fake_call_server import FakeCallServer
from .fake_incoming import FakeIncomingCaller
from .fake_server import FakeRegistrarServer

GOOD_PASSWORD = "clave-buena"
WRONG_PASSWORD = "clave-mala"

_ROOT_HOLD: list[tk.Tk] = []


@unittest.skipIf(tk is None, "tkinter no disponible")
class MainWindowTest(unittest.TestCase):
    def setUp(self) -> None:
        self._appdata = tempfile.mkdtemp(prefix="sipip-test-")
        self._previous_appdata = os.environ.get("APPDATA")
        os.environ["APPDATA"] = self._appdata

        self.server = FakeRegistrarServer(password=GOOD_PASSWORD, expires=300)
        self.server.start()

        self.root = tk.Tk()
        self.root.withdraw()
        _ROOT_HOLD.append(self.root)
        self.tones = DummyTones()
        self.window = MainWindow(self.root, tones=self.tones)
        self.page = self.window.setup_page
        self.page.vars["keep_password"].set(False)
        self._fill(extension="100", password=WRONG_PASSWORD)
        self.root.update()

    def tearDown(self) -> None:
        try:
            if self.window.registrar is not None and self.window.registrar.running:
                self.window.registrar.stop(timeout=0.5)
            self.window._cancel_scheduled()
            self.window.dialpad_page.engine.close()
            self.window.store.close()
            self.root.withdraw()
        except tk.TclError:
            pass
        self.server.stop()
        try:
            storage.delete_password("127.0.0.1", "100")
        except storage.CredentialError:
            pass
        if self._previous_appdata is None:
            os.environ.pop("APPDATA", None)
        else:
            os.environ["APPDATA"] = self._previous_appdata
        shutil.rmtree(self._appdata, ignore_errors=True)

    def _fill(self, extension: str, password: str, server: str | None = None) -> None:
        self.page.vars["server_host"].set(server or self.server.host)
        self.page.vars["server_port"].set(str(self.server.port))
        self.page.vars["extension"].set(extension)
        self.page.vars["password"].set(password)

    def _pump(self, seconds: float = 0.4) -> None:
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.01)

    def _wait_for_state(self, expected: RegistrarState, timeout: float = 10.0) -> None:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            registrar = self.window.registrar
            if registrar is not None and registrar.state is expected:
                return
            self.root.update()
            time.sleep(0.02)
        self.fail(f"el estado nunca fue {expected}")

    def _button(self, name: str) -> str:
        return str(getattr(self.page, name)["state"])

    def _register_successfully(self) -> None:
        self._fill(extension="100", password=GOOD_PASSWORD)
        self.window.register()
        self._wait_for_state(RegistrarState.REGISTERED)

    def test_wrong_password_shows_error_and_enables_retry(self) -> None:
        self.window.register()
        self._wait_for_state(RegistrarState.FAILED)
        self._pump()

        self.assertIn("Forbidden", self.window.registrar.detail)
        self.assertEqual(self._button("register_button"), "normal")
        self.assertEqual(self._button("next_button"), "disabled")

    def test_retry_after_failure_registers(self) -> None:
        self.window.register()
        self._wait_for_state(RegistrarState.FAILED)
        self._pump()

        self.page.vars["password"].set(GOOD_PASSWORD)
        self.window.register()
        self._wait_for_state(RegistrarState.REGISTERED)
        self._pump()

        self.assertEqual(self._button("register_button"), "disabled")
        self.assertEqual(self._button("unregister_button"), "normal")
        self.assertGreaterEqual(self.server.registrations, 1)

    def test_next_button_opens_dialpad(self) -> None:
        self._register_successfully()
        self._pump()

        self.assertEqual(self._button("next_button"), "normal")
        self.window.setup_page.next_button.invoke()
        self.root.update()

        self.assertEqual(self.window._current_page, "dialpad")

    def test_dialpad_requires_digits_and_registration(self) -> None:
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page

        self.assertEqual(str(dialpad.call_button["state"]), "disabled")

        for digit in "100":
            dialpad._append_digit(digit)
        self.root.update()
        self.assertEqual(dialpad.call_button["state"], "disabled")

        self._register_successfully()
        self._pump()
        self.assertEqual(dialpad.call_button["state"], "normal")
        self.assertTrue(dialpad.hint.cget("text").startswith("-> 100"))

    def test_dialpad_shows_prefix_and_contact_name(self) -> None:
        self.window.store.contacts.add("Ana", "100", favorite=True)
        self.page.vars["outbound_prefix"].set("8")
        self._register_successfully()
        self.assertEqual(self.window.account.outbound_prefix, "8")
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page

        for digit in "100":
            dialpad._append_digit(digit)
        self.root.update()
        self.assertIn("8 + 100", dialpad.hint.cget("text"))
        self.assertIn("Ana", dialpad.hint.cget("text"))

    def test_dialpad_backspace_and_clear(self) -> None:
        self._register_successfully()
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page

        for digit in "100":
            dialpad._append_digit(digit)
        dialpad._backspace()
        self.assertEqual(dialpad._digits, "10")
        dialpad._append_plus()
        self.assertEqual(dialpad._digits, "+10")
        dialpad._clear_digits()
        self.assertEqual(dialpad._digits, "")

    def test_dialing_starts_and_hangs_up(self) -> None:
        class _StubCallEngine:
            """Sustituye al motor real: determinista y sin hilos."""

            def __init__(self) -> None:
                self.session = CallSession()

            @property
            def state(self) -> CallState:
                return self.session.state

            def start_call(self, target, account, password) -> CallSession:
                self.session = CallSession(
                    target=target, peer=target, state=CallState.DIALING,
                    detail="Llamando...", started_at=time.time(),
                )
                return self.session

            def hangup(self) -> None:
                if self.state is not CallState.IDLE:
                    self.session = replace(
                        self.session, state=CallState.ENDED, detail="Llamada finalizada"
                    )

            def set_mute(self, enabled: bool) -> None:
                self.session = replace(self.session, muting=enabled)

            def set_hold(self, enabled: bool) -> None:
                self.session = replace(self.session, on_hold=enabled)

            def close(self) -> None:
                pass

        self._register_successfully()
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page
        dialpad.engine = _StubCallEngine()
        for digit in "100":
            dialpad._append_digit(digit)

        dialpad._toggle_call()
        self.root.update()
        self.assertEqual(dialpad.engine.state, CallState.DIALING)
        self.assertEqual(dialpad.call_state.cget("text"), "Marcando")

        dialpad._toggle_call()
        self.root.update()
        self.assertEqual(dialpad.engine.state, CallState.ENDED)

    def test_ended_call_shows_summary_and_exports_xlsx(self) -> None:
        from openpyxl import load_workbook

        class _StubCallEngine:
            def __init__(self) -> None:
                self.session = CallSession()

            @property
            def state(self) -> CallState:
                return self.session.state

            def start_call(self, target, account, password) -> CallSession:
                self.session = CallSession(
                    target=target, peer=target, state=CallState.DIALING,
                    detail="Llamando...", started_at=time.time(),
                )
                return self.session

            def hangup(self) -> None:
                pass

            def set_mute(self, enabled: bool) -> None:
                pass

            def set_hold(self, enabled: bool) -> None:
                pass

            def close(self) -> None:
                pass

        self._register_successfully()
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page
        dialpad.engine = _StubCallEngine()
        for digit in "6500":
            dialpad._append_digit(digit)
        dialpad._toggle_call()
        self.root.update()
        self.assertIsNotNone(dialpad._call_log_id)

        connected = CallSession(
            target="6500", peer="6500", state=CallState.CONNECTED,
            answered_at=time.time(), detail="En conversacion",
        )
        dialpad._engine_event_tap(connected)
        self._pump(0.3)
        self.assertIsNotNone(dialpad._answered_at)

        dialpad._engine_event_tap(
            replace(connected, state=CallState.ENDED, detail="Llamada finalizada")
        )
        self._pump(0.3)

        window = dialpad._summary_window
        self.assertIsNotNone(window)
        self.assertTrue(window.winfo_exists())
        values = dict(dialpad._summary_fields)
        self.assertEqual(values["Destino"], "6500")
        self.assertEqual(values["Estado"], "Completada")
        self.assertIn("C$", values["Monto"])

        path = os.path.join(tempfile.gettempdir(), "resumen_6500_test.xlsx")
        with patch(
            "sipclient.ui.dialpad_page.filedialog.asksaveasfilename", return_value=path
        ), patch("sipclient.ui.dialpad_page.messagebox.showinfo"):
            dialpad._export_summary()

        workbook = load_workbook(path)
        try:
            sheet = workbook["Llamada"]
            rows = {
                sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
                for row in range(1, sheet.max_row + 1)
            }
            self.assertEqual(rows["Campo"], "Valor")
            self.assertEqual(rows["Destino"], "6500")
            self.assertIn("C$", rows["Monto"])
        finally:
            workbook.close()
        os.remove(path)

    def test_long_press_zero_adds_plus(self) -> None:
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page
        dialpad._append_digit("6")
        dialpad._append_plus()
        self.assertTrue(dialpad._digits.startswith("+"))

    def test_clear_credentials_keeps_server_and_extension(self) -> None:
        self._register_successfully()
        self.window.clear_credentials()
        self._pump(0.2)

        self.assertEqual(self.page.vars["server_host"].get(), self.server.host)
        self.assertEqual(self.page.vars["extension"].get(), "100")
        self.assertEqual(self.page.vars["password"].get(), "")
        self.assertEqual(self._button("next_button"), "disabled")
        self.assertEqual(self._button("register_button"), "normal")
        self.assertTrue(
            self.server.wait_for_deregistrations(1, timeout=2.0),
            "no se libero el contacto en el servidor",
        )

    def test_navigation_updates_sidebar(self) -> None:
        for key in ("dialpad", "history", "contacts", "settings", "dialpad"):
            self.window.navigate(key)
            self.root.update()
            self.assertEqual(self.window._current_page, key)
            active = [
                name
                for name, button in self.window.sidebar_buttons.items()
                if str(button.cget("bg")) == SIDEBAR_ACTIVE
            ]
            self.assertEqual(active, [key])

    def test_contacts_roundtrip_and_dial(self) -> None:
        self.window.navigate("contacts")
        contacts_page = self.window.contacts_page

        contacts_page.name_var.set("Ana")
        contacts_page.number_var.set("100")
        contacts_page.favorite_var.set(True)
        contacts_page._save()
        self.root.update()
        self.assertEqual(len(contacts_page.contacts), 1)

        self._register_successfully()
        self.window.dial("100")
        self.root.update()
        self.assertEqual(self.window._current_page, "dialpad")
        self.assertEqual(self.window.dialpad_page._digits, "100")
        self.assertIn("Ana", self.window.dialpad_page.hint.cget("text"))

        self.window.navigate("dialpad")
        self.assertTrue(self.window.dialpad_page.favorites_frame.winfo_children())

    def test_history_lists_records_and_filters(self) -> None:
        self.window.store.calls.add(DIRECTION_OUTGOING, "100", STATUS_COMPLETED, duration=42, started_at=time.time())
        self.window.store.calls.add(DIRECTION_OUTGOING, "200", STATUS_COMPLETED, duration=10, started_at=time.time() - 60)

        self.window.navigate("history")
        history = self.window.history_page
        self.assertEqual(len(history.records), 2)

        history.search.set_value("200")
        self.root.update()
        self.assertEqual([record.number for record in history.records], ["200"])

        history.search.set_value("")
        history._set_filter("Recibidas")
        self.assertEqual(history.records, [])

        history._set_filter("Realizadas")
        self.assertEqual(len(history.records), 2)

    def test_voicemail_badge_from_server(self) -> None:
        self.server.voicemail = 2
        self._register_successfully()
        self._pump(0.3)
        self.assertEqual(self.window.registrar.unread_voicemail(), 2)
        self.assertIn("2", self.window.voicemail_badge.cget("text"))

    def test_status_bar_shows_refresh_countdown(self) -> None:
        self._register_successfully()
        self.window._tick()
        self.root.update()
        self.assertIn("refresco en", self.window.status_account.cget("text"))

    def test_pages_fill_the_content_area(self) -> None:
        self._map_window()
        sizes = set()
        for key in ("settings", "dialpad", "history", "contacts"):
            self.window.navigate(key)
            self.root.update_idletasks()
            frame = self.window.pages[key]
            self.assertGreater(frame.winfo_width(), 100)
            sizes.add((frame.winfo_width(), frame.winfo_height()))
            self.assertEqual(frame.winfo_width(), self.window.pages_frame.winfo_width())
            self.assertEqual(frame.winfo_height(), self.window.pages_frame.winfo_height())
        self.assertEqual(len(sizes), 1)

    def test_layout_has_no_overflow_with_long_text(self) -> None:
        long_text = (
            "Credenciales rechazadas por el servidor Asterisk de la centralita interna "
            "porque el prefijo de autorizacion no coincide con la extension configurada."
        )
        self._map_window()
        dialpad = self.window.dialpad_page
        dialpad.hint.configure(text=long_text)
        dialpad.call_detail.configure(text=long_text)
        self.window.pages["settings"].session_label.configure(text=long_text)
        self.window.log(long_text)

        problems: list[str] = []
        checked = 0
        for key in ("settings", "dialpad", "history", "contacts"):
            self.window.navigate(key)
            self.root.update_idletasks()
            checked += self._check_geometry(self.window.pages[key], key, problems)
        self.assertGreater(checked, 40)
        self.assertEqual(problems, [])
        self.root.withdraw()

    def _map_window(self) -> None:
        self.root.deiconify()
        self.root.geometry("980x700")
        self.root.update()
        self.root.update_idletasks()
        self.assertGreater(self.root.winfo_width(), 500)

    def _check_geometry(self, widget: tk.Misc, path: str, problems: list[str]) -> int:
        checked = 0
        for child in widget.winfo_children():
            if not child.winfo_ismapped():
                continue
            checked += 1
            name = path + "/" + child.__class__.__name__
            width, height = child.winfo_width(), child.winfo_height()
            if width <= 1 or height <= 1:
                problems.append(f"{name} degenerado {width}x{height}")
            if (
                child.winfo_x() + width > widget.winfo_width() + 2
                or child.winfo_y() + height > widget.winfo_height() + 2
            ):
                problems.append(f"{name} desborda {width}x{height}")
            checked += self._check_geometry(child, name, problems)
        return checked

    def test_long_contact_name_stays_inside_the_card(self) -> None:
        self.window.store.contacts.add(
            "Consola central de supervision muy larga", "900", favorite=True
        )
        self._map_window()
        self.window.navigate("dialpad")
        self.window.dialpad_page._set_digits("900")
        self.root.update_idletasks()
        problems: list[str] = []
        self._check_geometry(self.window.pages["dialpad"], "dialpad", problems)
        self.assertEqual(problems, [])
        self.root.withdraw()

    def _pump_until(self, predicate, timeout: float = 8.0) -> None:
        errors: list[BaseException] = []
        previous = threading.excepthook

        def _hook(args: threading.ExceptHookArgs) -> None:
            errors.append(args.exc_value)

        threading.excepthook = _hook
        try:
            deadline = time.monotonic() + timeout
            while time.monotonic() < deadline:
                self.root.update()
                if predicate():
                    return
                time.sleep(0.02)
        finally:
            threading.excepthook = previous
        if errors:
            raise errors[0]
        self.fail("la condicion no se cumplio a tiempo")

    def _deliver_incoming(self) -> FakeIncomingCaller:
        dialpad = self.window.dialpad_page
        caller = FakeIncomingCaller()
        caller.start()
        caller.call("100", ("127.0.0.1", dialpad.engine.listening_port))
        self._pump_until(lambda: str(dialpad.call_button["text"]) == "Contestar")
        return caller

    def test_incoming_call_answered_and_terminated(self) -> None:
        self._register_successfully()
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page

        caller = self._deliver_incoming()
        self.assertIn("entrante", dialpad.call_state.cget("text").lower())
        self.assertIn(("llamada_entrante", True), self.tones.played)

        dialpad.answer_button.invoke()
        self._pump_until(
            lambda: caller.established and dialpad.engine.session.state is CallState.CONNECTED
        )
        self.assertFalse(self.tones.playing, "el tono entrante debia detenerse")

        caller.bye()
        self._pump_until(lambda: dialpad.engine.session.state is CallState.ENDED)
        self._pump_until(
            lambda: any(
                r.direction == DIRECTION_INCOMING and r.status == STATUS_COMPLETED
                for r in self.window.store.calls.list()
            )
        )
        records = self.window.store.calls.list()
        self.assertTrue(
            any(r.direction == DIRECTION_INCOMING and r.status == STATUS_COMPLETED
                for r in records)
        )
        self.assertEqual(str(dialpad.call_button["state"]), "normal")

    def test_incoming_call_rejected(self) -> None:
        self._register_successfully()
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page

        caller = self._deliver_incoming()
        self.assertIn(("llamada_entrante", True), self.tones.played)

        dialpad.reject_button.invoke()
        self._pump_until(
            lambda: caller.rejected and dialpad.engine.session.state is CallState.ENDED
        )
        self.assertIn(486, caller.statuses)
        self._pump_until(
            lambda: any(
                r.direction == DIRECTION_INCOMING and r.status == STATUS_REJECTED
                for r in self.window.store.calls.list()
            )
        )

        records = self.window.store.calls.list()
        self.assertTrue(
            any(r.direction == DIRECTION_INCOMING and r.status == STATUS_REJECTED
                for r in records)
        )
        self.assertEqual(str(dialpad.call_button["text"]), "Llamar")
        self.assertNotIn(("telefono_colgado", False), self.tones.played)

    def test_no_answer_plays_tone_and_ends_call(self) -> None:
        self._register_successfully()
        self.window.navigate("dialpad")
        dialpad = self.window.dialpad_page

        call_server = FakeCallServer(password="secreto123", ringing=True, answer=False)
        call_server.start()
        self.addCleanup(call_server.stop)

        approach = dialpad_page.NO_ANSWER_TIMEOUT
        self.addCleanup(setattr, dialpad_page, "NO_ANSWER_TIMEOUT", approach)
        dialpad_page.NO_ANSWER_TIMEOUT = 0.5

        dialpad.engine.start_call(
            "200",
            AccountConfig(
                server_host=call_server.host,
                server_port=call_server.port,
                extension="100",
                domain="127.0.0.1",
            ),
            "secreto123",
        )

        self._pump_until(
            lambda: dialpad.engine.session.state is CallState.ENDED
            and call_server.cancels >= 1
            and ("fuera_de_servicio", False) in self.tones.played,
            timeout=10.0,
        )
        self.assertNotIn(("telefono_colgado", False), self.tones.played)
        self._pump_until(
            lambda: any(
                r.direction == DIRECTION_OUTGOING and r.status == STATUS_NO_ANSWER
                for r in self.window.store.calls.list()
            )
        )
        self.assertEqual(str(dialpad.call_button["text"]), "Llamar")


if __name__ == "__main__":
    unittest.main()
