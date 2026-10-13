"""Pruebas del exportador a Excel (.xlsx) para el resumen de llamadas."""

from __future__ import annotations

import os
import tempfile
import unittest

from sipclient.export_excel import write_call_summary


class WriteCallSummaryTest(unittest.TestCase):
    def test_writes_header_and_values(self) -> None:
        path = os.path.join(tempfile.mkdtemp(), "resumen.xlsx")
        write_call_summary(
            path,
            [
                ("Estado", "Completada"),
                ("Destino", "6500"),
                ("Duracion", "00:30 (30 s)"),
                ("Monto", "C$7.33"),
            ],
        )
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        try:
            sheet = workbook["Llamada"]
            self.assertEqual(sheet["A1"].value, "Campo")
            self.assertEqual(sheet["B1"].value, "Valor")
            rows = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                rows.append(tuple(row))
            self.assertEqual(
                rows,
                [
                    ("Estado", "Completada"),
                    ("Destino", "6500"),
                    ("Duracion", "00:30 (30 s)"),
                    ("Monto", "C$7.33"),
                ],
            )
        finally:
            workbook.close()

    def test_freeze_panes_and_arrow_down(self) -> None:
        path = os.path.join(tempfile.mkdtemp(), "resumen2.xlsx")
        write_call_summary(path, [("Destino", "100")])
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        try:
            sheet = workbook["Llamada"]
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet["A2"].value, "Destino")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()