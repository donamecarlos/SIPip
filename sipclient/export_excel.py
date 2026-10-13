"""Exportacion a Excel (.xlsx) del resumen de una llamada."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def write_call_summary(path: str | Path, fields: list[tuple[str, str]]) -> None:
    """Guarda el resumen como un libro con dos columnas: Campo y Valor."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Llamada"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    label_font = Font(bold=True)
    edge = Side(style="thin", color="D0D0D0")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)

    sheet["A1"] = "Campo"
    sheet["B1"] = "Valor"
    for cell in (sheet["A1"], sheet["B1"]):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for index, (label, value) in enumerate(fields, start=2):
        cell_label = sheet.cell(row=index, column=1, value=label)
        cell_label.font = label_font
        cell_label.border = border
        cell_value = sheet.cell(row=index, column=2, value=value)
        cell_value.border = border

    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 44
    workbook.save(path)